#!/usr/bin/env python3
"""Залить матрицу фурнитуры (BOM) из Excel прямо в Supabase через REST API.

Использует anon-ключ из fronted/.env.local. RLS-политики таблиц hardware_*
разрешают insert для anon. Запускать один раз после применения миграций со схемой.
"""
from __future__ import annotations

import json
import sys
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ENV_FILE = ROOT / "fronted" / ".env.local"
DEFAULT_XLSX = Path(r"C:\Users\ПК\Downloads\Шаблон ТЗ для заполнения.xlsx")

NAME_COL, SIZE_COL, SORT_COL = 4, 5, 2
FIRST_PRODUCT_COL, LAST_PRODUCT_COL = 6, 34

CONFIDENT_SECTIONS: dict[str, list[str]] = {
    "Стол Стабиле": ["Stabile"],
    "Стол примьер": ["Премьер черный", "Премьер белый"],
    "Авелла с ящиком": ["Avella"],
    "Авелла лайт": ["Avella lite"],
    "Донини гранде": ["Donini Grande 806", "Donini Grande 750"],
    "Донини с белыми ногами": ["Donini 806 белый", "Donini 750 белый"],
    "донини с черными ногами": ["Donini 806", "Donini 750"],
    "Классико +": ["Классико +"],
    "Классико": ["Классико"],
    "Кремона": ["Cremona"],
    "Донини Р": ["Donini R 806", "Donini R 750"],
    "Солито 2": ["Solito2"],
    "Солито": ["Solito 1150"],
    "Тумба тв 1500": ["ТВ Лофт 1500"],
    "Тумба тв 1100": ["ТВ Лофт"],
}


def read_env() -> tuple[str, str]:
    url = key = ""
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line.startswith("VITE_SUPABASE_URL="):
            url = line.split("=", 1)[1].strip()
        elif line.startswith("VITE_SUPABASE_ANON_KEY="):
            key = line.split("=", 1)[1].strip()
    if not url or not key:
        raise SystemExit("[error] no Supabase URL/anon key in .env.local")
    return url, key


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else ("%g" % v)
    return str(v).strip()


def pattern_from_name(product: str) -> str:
    base = product.strip()
    for token in ("Стол ", "Тумба ", "система ", "Система "):
        if base.startswith(token):
            base = base[len(token):]
    return "%" + base.strip() + "%"


def post(url: str, key: str, path: str, payload, prefer: str = "return=representation"):
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(f"{url}/rest/v1/{path}", data=data, method="POST")
    req.add_header("apikey", key)
    req.add_header("Authorization", f"Bearer {key}")
    req.add_header("Content-Type", "application/json")
    if prefer:
        req.add_header("Prefer", prefer)
    with urllib.request.urlopen(req, timeout=60) as resp:
        body = resp.read().decode("utf-8")
        return json.loads(body) if body else None


def rpc(url: str, key: str, name: str, args: dict):
    return post(url, key, f"rpc/{name}", args, prefer="")


def main() -> int:
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        raise SystemExit(f"[error] Excel not found: {xlsx_path}")
    url, key = read_env()

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    products: dict[int, str] = {}
    for col in range(FIRST_PRODUCT_COL, LAST_PRODUCT_COL + 1):
        name = norm(ws.cell(row=1, column=col).value)
        if name:
            products[col] = name

    items: list[dict] = []
    seen: set[tuple] = set()
    bom: dict[tuple, float] = defaultdict(float)

    for r in range(2, ws.max_row + 1):
        name = norm(ws.cell(row=r, column=NAME_COL).value)
        if not name:
            continue
        size = norm(ws.cell(row=r, column=SIZE_COL).value)
        sort_v = ws.cell(row=r, column=SORT_COL).value
        sort_order = int(sort_v) if isinstance(sort_v, (int, float)) else r
        key_ns = (name.lower(), size.lower())
        if key_ns not in seen:
            seen.add(key_ns)
            items.append({"name": name, "size": size, "sort_order": sort_order})
        for col, product in products.items():
            cell = ws.cell(row=r, column=col).value
            if isinstance(cell, (int, float)) and cell and cell > 0:
                bom[(name, size, product)] += float(cell)

    # 1) items -> get ids
    created = post(url, key, "hardware_items", items)
    id_by_key = {}
    for row in created or []:
        id_by_key[(str(row["name"]).strip().lower(), str(row.get("size") or "").strip().lower())] = row["id"]
    print(f"[ok] items inserted: {len(created or [])}")

    # 2) bom rows
    bom_payload = []
    skipped = 0
    for (name, size, product), qty in bom.items():
        item_id = id_by_key.get((name.strip().lower(), size.strip().lower()))
        if not item_id:
            skipped += 1
            continue
        bom_payload.append({"hardware_item_id": item_id, "bom_product": product, "qty_per_unit": qty})
    post(url, key, "hardware_bom", bom_payload)
    print(f"[ok] bom rows inserted: {len(bom_payload)} (skipped {skipped})")

    # 3) product map (via RPC, чтобы пройти CHECK и upsert)
    sort = 10
    mapped = 0
    for col in sorted(products):
        product = products[col]
        sections = CONFIDENT_SECTIONS.get(product)
        if sections:
            for sec in sections:
                rpc(url, key, "web_upsert_hardware_product_map_row", {
                    "p_id": None, "p_bom_product": product, "p_section_name": sec,
                    "p_item_name_pattern": None, "p_sort_order": sort, "p_is_active": True,
                })
                sort += 1
                mapped += 1
        else:
            rpc(url, key, "web_upsert_hardware_product_map_row", {
                "p_id": None, "p_bom_product": product, "p_section_name": None,
                "p_item_name_pattern": pattern_from_name(product), "p_sort_order": sort, "p_is_active": True,
            })
            sort += 1
            mapped += 1
    print(f"[ok] product map rows: {mapped}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
