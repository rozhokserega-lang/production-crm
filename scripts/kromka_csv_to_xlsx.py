#!/usr/bin/env python3
"""Convert kromka CSV export to Excel with proper columns."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "kromka_done_2026-06-01_2026-06-12.csv"
XLSX_PATH = ROOT / "kromka_done_2026-06-01_2026-06-12.xlsx"

WIDTHS = [12, 11, 9, 10, 42, 18, 8, 9, 14, 20, 14, 22, 22]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Кромка 1-12 июня"

    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            ws.append(row)

    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    wb.save(XLSX_PATH)
    print(f"Written {ws.max_row - 1} rows -> {XLSX_PATH}")


if __name__ == "__main__":
    main()
