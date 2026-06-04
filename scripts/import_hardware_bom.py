#!/usr/bin/env python3
"""Импорт матрицы соответствия фурнитуры (BOM) из Excel в seed-миграцию Supabase.

Excel «Шаблон ТЗ для заполнения.xlsx»:
  * строки  — позиции фурнитуры: D = наименование, E = размер, B = сортировка;
  * колонки F..AH — изделия (продукция);
  * ячейка — кол-во фурнитуры на 1 изделие.

Скрипт читает книгу и генерирует идемпотентную SQL-миграцию, наполняющую
hardware_items, hardware_bom и дефолтные строки hardware_product_map.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(r"C:\Users\ПК\Downloads\Шаблон ТЗ для заполнения.xlsx")
OUT_SQL = ROOT / "supabase" / "migrations" / "20260604140200_hardware_bom_seed.sql"

NAME_COL = 4   # D — наименование фурнитуры
SIZE_COL = 5   # E — размер
SORT_COL = 2   # B — сортировка
FIRST_PRODUCT_COL = 6   # F
LAST_PRODUCT_COL = 34   # AH

# Достоверные сопоставления изделия из BOM -> секция(и) плана (section_catalog).
# Остальные изделия получают item_name_pattern по очищенному названию и
# правятся вручную в админке («Соответствие»).
CONFIDENT_SECTIONS: dict[str, list[str]] = {
    "Стол Стабиле": ["Stabile"],
    "Стол примьер": ["Премьер черный", "Премьер белый"],
    "Авелла с ящиком": ["Avella"],
    "Авелла лайт": ["Avella lite"],
    "Донини гранде": ["Donini Grande 806", "Donini Grande 750"],
    "Донини с белыми ногами": ["Donini 806 белый", "Donini 750 белый"],
    "донини с черными ногами": ["Donini 806", "Donini 750"],
    "Классико +": ["Классико +"],
    "Классико": ["Классико"],
    "Кремона": ["Cremona"],
    "Донини Р": ["Donini R 806", "Donini R 750"],
    "Солито 2": ["Solito2"],
    "Солито": ["Solito 1150"],
    "Тумба тв 1500": ["ТВ Лофт 1500"],
    "Тумба тв 1100": ["ТВ Лофт"],
}


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ("%g" % v)
    return str(v).strip()


def sql_str(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def pattern_from_name(product: str) -> str:
    """Грубый ILIKE-паттерн для неуверенных сопоставлений."""
    base = product.strip()
    for token in ("Стол ", "Тумба ", "система ", "Система "):
        if base.startswith(token):
            base = base[len(token):]
    base = base.strip()
    return "%" + base + "%"


def main() -> int:
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        print(f"[error] Excel not found: {xlsx_path}")
        return 1

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    products: dict[int, str] = {}
    for col in range(FIRST_PRODUCT_COL, LAST_PRODUCT_COL + 1):
        name = norm(ws.cell(row=1, column=col).value)
        if name:
            products[col] = name

    items: list[dict] = []           # {name, size, sort}
    bom: dict[tuple, float] = defaultdict(float)  # (name, size, product) -> qty
    seen_items: set[tuple] = set()

    for r in range(2, ws.max_row + 1):
        name = norm(ws.cell(row=r, column=NAME_COL).value)
        if not name:
            continue
        size = norm(ws.cell(row=r, column=SIZE_COL).value)
        sort_v = ws.cell(row=r, column=SORT_COL).value
        sort_order = int(sort_v) if isinstance(sort_v, (int, float)) else r

        key = (name.lower(), size.lower())
        if key not in seen_items:
            seen_items.add(key)
            items.append({"name": name, "size": size, "sort": sort_order})

        for col, product in products.items():
            cell = ws.cell(row=r, column=col).value
            if isinstance(cell, (int, float)) and cell:
                qty = float(cell)
                if qty > 0:
                    bom[(name, size, product)] += qty

    # ---- emit SQL ----
    lines: list[str] = []
    lines.append("-- ============================================================")
    lines.append("-- SEED: матрица соответствия фурнитуры (BOM).")
    lines.append(f"-- Сгенерировано scripts/import_hardware_bom.py из {xlsx_path.name}")
    lines.append("-- Идемпотентно (ON CONFLICT). Правьте соответствия в админке.")
    lines.append("-- ============================================================")
    lines.append("")

    lines.append("-- 1) Позиции фурнитуры")
    for it in items:
        lines.append(
            "INSERT INTO public.hardware_items (name, size, sort_order) VALUES ("
            f"{sql_str(it['name'])}, {sql_str(it['size'])}, {it['sort']})"
            " ON CONFLICT (lower(trim(name)), lower(trim(coalesce(size, '')))) "
            "DO UPDATE SET sort_order = EXCLUDED.sort_order, is_active = TRUE;"
        )
    lines.append("")

    lines.append("-- 2) Нормы расхода (фурнитура x изделие)")
    for (name, size, product), qty in sorted(bom.items(), key=lambda kv: (kv[0][2], kv[0][0])):
        qty_s = ("%g" % qty)
        lines.append(
            "INSERT INTO public.hardware_bom (hardware_item_id, bom_product, qty_per_unit) "
            f"SELECT i.id, {sql_str(product)}, {qty_s} FROM public.hardware_items i "
            f"WHERE lower(trim(i.name)) = lower(trim({sql_str(name)})) "
            f"AND lower(trim(coalesce(i.size, ''))) = lower(trim({sql_str(size)})) "
            "ON CONFLICT (hardware_item_id, lower(trim(bom_product))) "
            "DO UPDATE SET qty_per_unit = EXCLUDED.qty_per_unit;"
        )
    lines.append("")

    lines.append("-- 3) Соответствие изделий BOM -> секции плана (дефолт, правится в админке)")
    sort = 10
    for col in sorted(products):
        product = products[col]
        sections = CONFIDENT_SECTIONS.get(product)
        if sections:
            for sec in sections:
                lines.append(
                    "INSERT INTO public.hardware_product_map (bom_product, section_name, sort_order) VALUES ("
                    f"{sql_str(product)}, {sql_str(sec)}, {sort})"
                    " ON CONFLICT (lower(trim(bom_product)), coalesce(lower(trim(section_name)), ''), "
                    "coalesce(lower(trim(item_name_pattern)), '')) DO NOTHING;"
                )
                sort += 1
        else:
            lines.append(
                "INSERT INTO public.hardware_product_map (bom_product, item_name_pattern, sort_order) VALUES ("
                f"{sql_str(product)}, {sql_str(pattern_from_name(product))}, {sort})"
                " ON CONFLICT (lower(trim(bom_product)), coalesce(lower(trim(section_name)), ''), "
                "coalesce(lower(trim(item_name_pattern)), '')) DO NOTHING;"
            )
            sort += 1
    lines.append("")

    OUT_SQL.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"[ok] items={len(items)} bom_rows={len(bom)} products={len(products)}")
    print(f"[ok] wrote {OUT_SQL}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
